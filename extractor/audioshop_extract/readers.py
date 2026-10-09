"""One reader per kind of input. Each returns a Document: a title guess, the body text with blank
lines between paragraphs (single newlines are line wrapping), and warnings worth showing the reader.
"""

from __future__ import annotations

import io
import re
import zipfile
from collections import Counter
from dataclasses import dataclass, field
from pathlib import Path, PurePosixPath

from audioshop_extract import netguard, ocr

IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".webp", ".heic", ".heif", ".gif", ".bmp", ".tif", ".tiff"}
TEXT_EXTS = {".txt", ".text", ".md", ".markdown"}


class ExtractError(Exception):
    """A user-facing reason the text could not be extracted."""


@dataclass
class Document:
    title: str = ""
    text: str = ""
    kind: str = ""
    pages: int = 1
    warnings: list[str] = field(default_factory=list)


# ------------------------------------------------------------------------------ text helpers


def tidy(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("­", "").replace(" ", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r" *\n *", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def decode(data: bytes) -> str:
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def strip_markdown(md: str) -> str:
    md = re.sub(r"```.*?```", "", md, flags=re.S)
    md = re.sub(r"!\[[^\]]*\]\([^)]*\)", "", md)
    md = re.sub(r"\[([^\]]+)\]\([^)]*\)", r"\1", md)
    md = re.sub(r"^\s{0,3}#{1,6}\s*", "", md, flags=re.M)
    md = re.sub(r"^\s{0,3}>\s?", "", md, flags=re.M)
    md = re.sub(r"^\s*[-*+]\s+", "", md, flags=re.M)
    md = re.sub(r"(\*\*|__|\*|_|`)(?=\S)(.+?)(?<=\S)\1", r"\2", md)
    return md


def first_line_title(text: str) -> str:
    first = text.strip().split("\n", 1)[0].strip()
    if 0 < len(first) <= 100 and not first.endswith((".", "!", "?", ",", ";", ":")):
        return first
    return ""


def drop_leading_title(text: str, title: str) -> str:
    """Remove the title from the start of the body (it is announced separately)."""
    body = text.lstrip()
    if title and body[: len(title)].strip().lower() == title.strip().lower():
        return body[len(title):].lstrip(" \n:-–—")
    return text


# --------------------------------------------------------------------------------- readers


def read_text(data: bytes, name: str) -> Document:
    raw = decode(data)
    if Path(name).suffix.lower() in (".md", ".markdown"):
        raw = strip_markdown(raw)
    text = tidy(raw)
    return Document(first_line_title(text), text, "text")


def read_html(data: bytes, url: str | None = None) -> Document:
    import trafilatura

    text = trafilatura.extract(data, url=url, favor_precision=True, include_comments=False,
                               include_tables=False, include_images=False, include_links=False,
                               deduplicate=True)
    if not text:
        raise ExtractError("No article text found on that page. It may need a login or be mostly "
                           "pictures or video; paste the text instead.")
    meta = trafilatura.extract_metadata(data)
    title = meta.title.strip() if meta and meta.title else ""
    text = "\n\n".join(line.strip() for line in text.splitlines() if line.strip())
    return Document(title, tidy(text), "web page")


def read_pdf(data: bytes) -> Document:
    from pypdf import PdfReader

    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            reader.decrypt("")
    except Exception as exc:
        raise ExtractError("That PDF could not be opened (damaged or password-protected).") from exc
    pages, scanned = [], []
    for i, page in enumerate(reader.pages):
        try:
            text = page.extract_text() or ""
        except Exception:
            text = ""
        if len(text.strip()) < 30:
            scanned.append(i)
        pages.append(text)
    warnings = []
    if scanned:
        import pypdfium2

        pdf = pypdfium2.PdfDocument(data)
        for i in scanned:
            w, h = pdf[i].get_size()  # points; aim for ~200 dpi but never past ocr.MAX_SIDE pixels
            img = pdf[i].render(scale=min(200 / 72, ocr.MAX_SIDE / max(w, h))).to_pil().convert("RGB")
            text, conf, cols = ocr.ocr_image(img)
            pages[i] = text
            if text and conf < 0.8:
                warnings.append(f"Page {i + 1} was read from an image and may contain mistakes.")
        warnings.insert(0, f"{len(scanned)} page(s) had no text layer and were read with OCR.")
    text = _strip_running_lines(pages)
    meta_title = ""
    try:
        meta_title = (reader.metadata.title or "").strip() if reader.metadata else ""
    except Exception:
        pass
    return Document(meta_title or first_line_title(text), text, "PDF", len(pages), warnings)


def _strip_running_lines(pages: list[str]) -> str:
    """Join PDF pages, dropping page numbers and headers/footers repeated on most pages, and
    turning the PDF's hard line breaks into paragraphs."""
    split = [[l.strip() for l in p.splitlines() if l.strip()] for p in pages]
    norm = lambda l: re.sub(r"\d+", "#", l.lower())  # noqa: E731
    counts = Counter(norm(l) for lines in split for l in set(lines))
    # Headers and footers are short lines repeated on most pages (digits ignored: "Page 3 of 9").
    repeated = {k for k, c in counts.items() if len(pages) >= 3 and c >= 0.5 * len(pages) and len(k) < 70}
    out = []
    for lines in split:
        kept = [l for l in lines if norm(l) not in repeated
                and not re.fullmatch(r"(page\s*)?\d+(\s*(of|/)\s*\d+)?", l, re.I)]
        if not kept:
            continue
        width = sorted(len(l) for l in kept)[int(0.8 * (len(kept) - 1))]
        paragraph, paragraphs = [], []
        for line in kept:
            paragraph.append(line)
            if line.endswith((".", "!", "?", ":", "”", '"')) and len(line) < 0.8 * width:
                paragraphs.append(ocr._join(paragraph))
                paragraph = []
        if paragraph:
            paragraphs.append(ocr._join(paragraph))
        out.append("\n\n".join(paragraphs))
    return tidy("\n\n".join(out))


def read_docx(data: bytes) -> Document:
    import docx

    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:
        raise ExtractError("That Word file could not be opened.") from exc
    paragraphs = [p.text.strip() for p in document.paragraphs if p.text.strip()]
    title = (document.core_properties.title or "").strip()
    if not title:
        for p in document.paragraphs:
            if p.text.strip() and (p.style.name or "").lower().startswith(("title", "heading")):
                title = p.text.strip()
                break
    text = tidy("\n\n".join(paragraphs))
    return Document(title or first_line_title(text), text, "Word document")


def read_epub(data: bytes) -> Document:
    from lxml import etree, html as lxml_html

    try:
        book = zipfile.ZipFile(io.BytesIO(data))
        container = etree.fromstring(book.read("META-INF/container.xml"))
        opf_path = container.find(".//{*}rootfile").get("full-path")
        opf = etree.fromstring(book.read(opf_path))
    except Exception as exc:
        raise ExtractError("That EPUB could not be opened.") from exc
    base = PurePosixPath(opf_path).parent
    items = {i.get("id"): i.get("href") for i in opf.iterfind(".//{*}manifest/{*}item")}
    spine = [items[r.get("idref")] for r in opf.iterfind(".//{*}spine/{*}itemref") if r.get("idref") in items]
    title_el = opf.find(".//{*}metadata/{http://purl.org/dc/elements/1.1/}title")
    paragraphs = []
    for href in spine:
        try:
            tree = lxml_html.fromstring(book.read(str(base / href)))
        except Exception:
            continue
        for el in tree.iter("p", "h1", "h2", "h3", "h4", "li", "blockquote"):
            t = " ".join(el.text_content().split())
            if t:
                paragraphs.append(t)
    if not paragraphs:
        raise ExtractError("No readable text found in that EPUB (it may be DRM-protected).")
    title = title_el.text.strip() if title_el is not None and title_el.text else ""
    return Document(title, tidy("\n\n".join(paragraphs)), "EPUB book")


def read_images(paths: list[Path]) -> Document:
    texts, warnings, columns = [], [], 0
    for n, path in enumerate(paths, 1):
        try:
            img = ocr.load_image(path)
        except Exception as exc:
            raise ExtractError(f"Photo {n} could not be opened as an image.") from exc
        text, conf, cols = ocr.ocr_image(img)
        if not text:
            warnings.append(f"No text found in photo {n}.")
            continue
        columns += cols
        if conf < 0.8:
            warnings.append(f"Photo {n} was hard to read; some words may be wrong.")
        texts.append(text)
    if not texts:
        raise ExtractError("No text found in the photo. Try a sharper, well-lit picture taken straight on.")
    if columns:
        warnings.append("Some pages looked like two columns; they were read left column first.")
    text = tidy("\n\n".join(texts))
    title = first_line_title(text)
    return Document(title, text, "photo" if len(paths) == 1 else "photos", len(paths), warnings)


def read_url(url: str) -> Document:
    fetched = netguard.fetch(url)
    if "pdf" in fetched.content_type or fetched.body[:5] == b"%PDF-":
        doc = read_pdf(fetched.body)
    elif fetched.content_type.startswith("text/plain"):
        doc = read_text(fetched.body, "page.txt")
    else:
        doc = read_html(fetched.body, fetched.url)
    if len(doc.text.split()) < 80:
        doc.warnings.append("Only a little text was found. The page may be paywalled or load its "
                            "article with JavaScript; if so, paste the text instead.")
    return doc


def read_file(path: Path, name: str) -> Document:
    ext = Path(name).suffix.lower()
    data = path.read_bytes()
    if ext in TEXT_EXTS:
        return read_text(data, name)
    if ext in (".html", ".htm", ".xhtml"):
        return read_html(data)
    if ext == ".pdf" or data[:5] == b"%PDF-":
        return read_pdf(data)
    if ext == ".docx":
        return read_docx(data)
    if ext == ".epub":
        return read_epub(data)
    if ext in IMAGE_EXTS:
        return read_images([path])
    raise ExtractError(f"Audioshop cannot read {ext or 'that kind of'} files yet.")
