"""Extraction tests. Fixtures are generated here, so nothing personal or copyrighted is committed."""

import io
import json
import subprocess
import sys
import zipfile
from pathlib import Path

import pytest
from PIL import Image, ImageDraw, ImageFont

from audioshop_extract import netguard
from audioshop_extract.readers import (
    ExtractError,
    drop_leading_title,
    read_docx,
    read_epub,
    read_file,
    read_html,
    read_images,
    read_pdf,
    read_text,
)

ARTICLE = [
    "City Council Approves New Bike Lanes",
    "The city council voted on Tuesday to add twelve miles of protected bike lanes downtown.",
    "Construction will begin in the spring and should finish before the end of next year.",
]


def page_image(lines: list[str], size=(2200, 900), font_size=40) -> Image.Image:
    img = Image.new("RGB", size, "white")
    draw = ImageDraw.Draw(img)
    font = ImageFont.load_default(size=font_size)
    y = 60
    for i, line in enumerate(lines):
        draw.text((60, y), line, fill="black", font=ImageFont.load_default(size=font_size + 14) if i == 0 else font)
        y += font_size * 2 + (30 if i == 0 else 0)
    return img


def test_plain_and_markdown_text():
    doc = read_text("﻿Hello world\n\nThis is   a test.\r\n".encode(), "a.txt")
    assert doc.title == "Hello world" and "This is a test." in doc.text
    md = read_text(b"# Title\n\nSome **bold** [link](http://x.y) text.", "a.md")
    assert "Some bold link text." in md.text and "#" not in md.text


def test_html_main_text_and_title():
    body = " ".join(["This paragraph is the real article text and it is long enough."] * 5)
    html = (f"<html><head><title>Bike Lanes</title></head><body><nav>Home | Sports</nav>"
            f"<article><h1>Bike Lanes</h1><p>{body}</p><p>{body}</p></article><footer>Related</footer></body></html>")
    doc = read_html(html.encode())
    assert doc.title == "Bike Lanes" and "real article text" in doc.text and "Sports" not in doc.text


def test_docx():
    import docx

    d = docx.Document()
    d.add_heading(ARTICLE[0], 0)
    for line in ARTICLE[1:]:
        d.add_paragraph(line)
    buf = io.BytesIO()
    d.save(buf)
    doc = read_docx(buf.getvalue())
    assert doc.title == ARTICLE[0] and ARTICLE[2] in doc.text


def test_epub_in_spine_order():
    container = ('<?xml version="1.0"?><container xmlns="urn:oasis:names:tc:opendocument:xmlns:container">'
                 '<rootfiles><rootfile full-path="OEBPS/book.opf"/></rootfiles></container>')
    opf = ('<?xml version="1.0"?><package xmlns="http://www.idpf.org/2007/opf"><metadata '
           'xmlns:dc="http://purl.org/dc/elements/1.1/"><dc:title>A Small Book</dc:title></metadata>'
           '<manifest><item id="b" href="two.xhtml"/><item id="a" href="one.xhtml"/></manifest>'
           '<spine><itemref idref="a"/><itemref idref="b"/></spine></package>')
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("META-INF/container.xml", container)
        z.writestr("OEBPS/book.opf", opf)
        z.writestr("OEBPS/one.xhtml", "<html><body><h1>Chapter one</h1><p>First words.</p></body></html>")
        z.writestr("OEBPS/two.xhtml", "<html><body><p>Second words.</p></body></html>")
    doc = read_epub(buf.getvalue())
    assert doc.title == "A Small Book"
    assert doc.text.index("First words.") < doc.text.index("Second words.")


def test_text_pdf_and_page_numbers_dropped():
    # A minimal hand-written PDF with a text layer on 3 pages (page number at the bottom of each).
    pages = []
    words = {1: "Alpha", 2: "Bravo", 3: "Charlie"}
    for n in (1, 2, 3):
        stream = (f"BT /F1 12 Tf 72 740 Td (Bike Lane Weekly) Tj 0 -20 Td ({words[n]} paragraph of the "
                  f"report ends here.) Tj 0 -650 Td ({n}) Tj ET").encode()
        pages.append(stream)
    objs = ["<< /Type /Catalog /Pages 2 0 R >>",
            "<< /Type /Pages /Kids [" + " ".join(f"{4 + 2 * i} 0 R" for i in range(3)) + "] /Count 3 >>",
            "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
    for i, s in enumerate(pages):
        objs.append(f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents {5 + 2 * i} 0 R "
                    "/Resources << /Font << /F1 3 0 R >> >> >>")
        objs.append(f"<< /Length {len(s)} >>\nstream\n{s.decode()}\nendstream")
    out, offsets = io.BytesIO(), []
    out.write(b"%PDF-1.4\n")
    for i, o in enumerate(objs, 1):
        offsets.append(out.tell())
        out.write(f"{i} 0 obj\n{o}\nendobj\n".encode())
    xref = out.tell()
    out.write(f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode())
    for off in offsets:
        out.write(f"{off:010d} 00000 n \n".encode())
    out.write(f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode())
    doc = read_pdf(out.getvalue())
    assert doc.pages == 3
    for word in ("Alpha", "Bravo", "Charlie"):
        assert f"{word} paragraph of the report ends here." in doc.text
    assert "Bike Lane Weekly" not in doc.text  # running header removed
    assert not any(line.strip().isdigit() for line in doc.text.splitlines())  # page numbers removed


@pytest.mark.slow
def test_scanned_pdf_uses_ocr():
    buf = io.BytesIO()
    page_image(ARTICLE).save(buf, format="PDF")
    doc = read_pdf(buf.getvalue())
    assert any("OCR" in w for w in doc.warnings)
    assert "bike lanes" in doc.text.lower() and "spring" in doc.text.lower()


@pytest.mark.slow
def test_photo_of_a_page(tmp_path):
    path = tmp_path / "page.jpg"
    page_image(ARTICLE).save(path, quality=90)
    doc = read_images([path])
    words = doc.text.lower()
    assert "council" in words and "protected bike lanes" in words and "next year" in words
    assert doc.title.lower().startswith("city council")


@pytest.mark.slow
def test_tilted_photo_keeps_headline_and_paragraphs(tmp_path):
    img = Image.new("RGB", (1240, 1000), (250, 248, 240))
    draw = ImageDraw.Draw(img)
    draw.text((80, 90), "Library extends weekend hours", fill="black", font=ImageFont.load_default(size=54))
    body = ImageFont.load_default(size=36)
    y = 200
    for para in (["The public library will stay open until nine on", "Saturdays and Sundays starting next month."],
                 ["Officials said the change follows a survey of", "residents who asked for more study space."]):
        for line in para:
            draw.text((80, y), line, fill="black", font=body)
            y += 52
        y += 30
    path = tmp_path / "tilted.png"
    img.rotate(-1.5, expand=True, fillcolor=(200, 200, 200)).save(path)
    doc = read_images([path])
    assert doc.title == "Library extends weekend hours"
    paragraphs = doc.text.split("\n\n")
    assert len(paragraphs) == 3, paragraphs
    assert paragraphs[2].startswith("Officials said")


def test_unknown_type_is_a_friendly_error(tmp_path):
    f = tmp_path / "a.xyz"
    f.write_bytes(b"\x00\x01")
    with pytest.raises(ExtractError, match="cannot read"):
        read_file(f, "a.xyz")


def test_title_not_read_twice():
    assert drop_leading_title("Bike Lanes\n\nThe council voted.", "Bike Lanes") == "The council voted."


@pytest.mark.parametrize("url", ["http://127.0.0.1/", "http://localhost:8787/api", "http://169.254.169.254/latest",
                                 "http://10.0.0.5/", "http://100.100.100.100/", "ftp://example.com/a"])
def test_private_and_odd_urls_refused(url, monkeypatch):
    monkeypatch.delenv("AUDIOSHOP_ALLOW_PRIVATE_HOSTS", raising=False)
    with pytest.raises(netguard.FetchError):
        netguard.fetch(url)


def test_cli_writes_text_and_meta(tmp_path):
    (tmp_path / "job.json").write_text(json.dumps({"source": {"type": "text", "text": "\n\n".join(ARTICLE)}}))
    root = Path(__file__).resolve().parents[1]
    proc = subprocess.run([sys.executable, "-m", "audioshop_extract", str(tmp_path)], cwd=root,
                          capture_output=True, text=True)
    assert proc.returncode == 0, proc.stderr
    meta = json.loads((tmp_path / "extract.json").read_text())
    assert meta["title"] == ARTICLE[0] and meta["titleFound"] is True and meta["words"] > 20
    assert not (tmp_path / "text.txt").read_text().startswith(ARTICLE[0])


def test_cli_labels_untitled_text_without_announcing_it(tmp_path):
    text = "On Tuesday the council voted for twelve new bike lanes. Work starts in spring."
    (tmp_path / "job.json").write_text(json.dumps({"source": {"type": "text", "text": text}}))
    root = Path(__file__).resolve().parents[1]
    proc = subprocess.run([sys.executable, "-m", "audioshop_extract", str(tmp_path)], cwd=root,
                          capture_output=True, text=True)
    assert proc.returncode == 0, proc.stderr
    meta = json.loads((tmp_path / "extract.json").read_text())
    assert meta["title"] == "On Tuesday the council voted for twelve new bike lanes"
    assert meta["titleFound"] is False
    assert (tmp_path / "text.txt").read_text().startswith("On Tuesday")  # nothing dropped


def test_cli_reports_user_facing_error(tmp_path):
    (tmp_path / "job.json").write_text(json.dumps({"source": {"type": "url", "url": "http://127.0.0.1:9/"}}))
    root = Path(__file__).resolve().parents[1]
    proc = subprocess.run([sys.executable, "-m", "audioshop_extract", str(tmp_path)], cwd=root,
                          capture_output=True, text=True)
    assert proc.returncode == 2
    assert proc.stderr.strip().splitlines()[-1].startswith("AUDIOSHOP_ERROR: Refusing")
